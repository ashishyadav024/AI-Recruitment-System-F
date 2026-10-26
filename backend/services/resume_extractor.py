"""
Resume Text Extraction Service
Extracts clean, formatted text from PDF, DOCX, and TXT resume files with:
- Magic bytes content inspection
- PyMuPDF page-by-page text extraction with header/footer deduplication
- python-docx paragraph and table extraction
- Multi-encoding fallback for text files
- Whitespace normalization and non-printable character removal
- Scanned image detection and corruption handling
- File size (5MB max) and text quality validation (<200 char warning)
- Separate storage of original file for recruiter access
"""

import os
import re
import io
import time
import uuid
import logging
from typing import Dict, Any, Tuple, Optional
import pymupdf
import docx

logger = logging.getLogger(__name__)

MAX_FILE_SIZE_BYTES = 5 * 1024 * 1024  # 5 MB
MIN_VALID_TEXT_LENGTH = 200            # Minimum meaningful resume text length
MAX_SAFE_TEXT_LENGTH = 60000           # Cap for safe LLM processing

UPLOAD_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "uploads", "resumes")
os.makedirs(UPLOAD_DIR, exist_ok=True)


class ResumeExtractionError(Exception):
    """Custom exception for resume extraction errors with user-friendly messages"""
    def __init__(self, message: str, status_code: int = 400):
        super().__init__(message)
        self.message = message
        self.status_code = status_code


def detect_file_type(content: bytes, filename: str) -> str:
    """
    Detect file format based on both extension and magic bytes.
    Returns: 'pdf', 'docx', or 'txt'
    Raises: ResumeExtractionError if unsupported or mismatched
    """
    ext = os.path.splitext(filename.lower())[1]

    if not content or len(content) == 0:
        raise ResumeExtractionError("The uploaded file is empty. Please upload a valid resume.", status_code=400)

    # Magic Bytes Inspection
    is_pdf_magic = content.startswith(b"%PDF-")
    is_zip_magic = content.startswith(b"PK\x03\x04")

    if ext == ".pdf":
        if not is_pdf_magic:
            raise ResumeExtractionError("Invalid PDF file: Missing standard PDF header.", status_code=400)
        return "pdf"

    elif ext == ".docx":
        if not is_zip_magic:
            raise ResumeExtractionError("Invalid DOCX file: Missing standard Office OpenXML header.", status_code=400)
        return "docx"

    elif ext == ".txt":
        # Check that txt does not contain binary null bytes
        if b"\x00" in content[:1024]:
            raise ResumeExtractionError("The uploaded .txt file contains binary characters.", status_code=400)
        return "txt"

    else:
        raise ResumeExtractionError(
            f"Unsupported file type '{ext}'. Please upload a .pdf, .docx, or .txt resume.",
            status_code=400
        )


def clean_and_normalize_text(raw_text: str) -> str:
    """
    Cleans raw extracted text:
    - Normalizes newlines (\r\n -> \n)
    - Removes null bytes and non-printable control characters
    - Normalizes excessive spacing while preserving bullet points and line structure
    - Limits text to a safe size
    """
    if not raw_text:
        return ""

    # 1. Normalize line endings
    text = raw_text.replace("\r\n", "\n").replace("\r", "\n")

    # 2. Remove null bytes and non-printable control characters (keep \t, \n)
    text = re.sub(r'[\x00-\x08\x0B\x0C\x0E-\x1F\x7F]', '', text)

    # 3. Clean line by line
    cleaned_lines = []
    for line in text.split("\n"):
        # Replace multiple horizontal spaces/tabs with single space
        cleaned_line = re.sub(r'[ \t]+', ' ', line).strip()
        cleaned_lines.append(cleaned_line)

    # 4. Collapse 3+ consecutive newlines to 2
    joined = "\n".join(cleaned_lines)
    joined = re.sub(r'\n{3,}', '\n\n', joined).strip()

    # 5. Cap maximum safe length
    if len(joined) > MAX_SAFE_TEXT_LENGTH:
        joined = joined[:MAX_SAFE_TEXT_LENGTH] + "\n\n[Text truncated to safe limit...]"

    return joined


def remove_repeated_headers_footers(pages_text: list[str]) -> list[str]:
    """
    Detects and strips repeated header or footer lines across multiple PDF pages.
    """
    if len(pages_text) < 2:
        return pages_text

    # Split each page into lines
    pages_lines = [p.split("\n") for p in pages_text if p]
    if len(pages_lines) < 2:
        return pages_text

    # Check top line across pages
    first_lines = [lines[0].strip() for lines in pages_lines if lines]
    if len(first_lines) >= 2 and all(l == first_lines[0] and len(l) > 5 for l in first_lines):
        # Remove repeated header from subsequent pages
        for i in range(1, len(pages_lines)):
            if pages_lines[i] and pages_lines[i][0].strip() == first_lines[0]:
                pages_lines[i] = pages_lines[i][1:]

    # Check bottom line across pages
    last_lines = [lines[-1].strip() for lines in pages_lines if lines]
    if len(last_lines) >= 2 and all(l == last_lines[0] and len(l) > 3 for l in last_lines):
        # Remove repeated footer from subsequent pages
        for i in range(1, len(pages_lines)):
            if pages_lines[i] and pages_lines[i][-1].strip() == last_lines[0]:
                pages_lines[i] = pages_lines[i][:-1]

    return ["\n".join(lines) for lines in pages_lines]


def extract_from_pdf(content: bytes) -> str:
    """Extract clean text page by page from PDF using PyMuPDF (fitz)"""
    try:
        doc = pymupdf.open(stream=content, filetype="pdf")
    except Exception as e:
        logger.error(f"Failed to open PDF: {e}")
        raise ResumeExtractionError("Corrupted PDF file. Unable to read document structure.", status_code=400)

    try:
        # Check encryption / password protection
        if doc.is_encrypted:
            raise ResumeExtractionError(
                "The uploaded PDF is password-protected. Please upload an unlocked PDF.",
                status_code=400
            )

        if len(doc) == 0:
            raise ResumeExtractionError("The uploaded PDF document contains no pages.", status_code=400)

        pages_text = []
        has_images = False

        for page in doc:
            page_text = page.get_text("text")
            if page_text and page_text.strip():
                pages_text.append(page_text.strip())
            if not has_images and len(page.get_images()) > 0:
                has_images = True

        # Header/footer deduplication
        deduped_pages = remove_repeated_headers_footers(pages_text)
        full_text = "\n\n".join(p for p in deduped_pages if p.strip())

        # Check for scanned image PDF
        if len(full_text.strip()) < 50:
            if has_images or len(doc) > 0:
                raise ResumeExtractionError(
                    "This PDF looks like a scanned image. Please upload a text-based PDF or paste your resume text.",
                    status_code=400
                )

        return clean_and_normalize_text(full_text)
    finally:
        doc.close()


def extract_from_docx(content: bytes) -> str:
    """Extract clean text from DOCX (paragraphs and tables) using python-docx"""
    try:
        doc = docx.Document(io.BytesIO(content))
    except Exception as e:
        logger.error(f"Failed to open DOCX: {e}")
        raise ResumeExtractionError("Corrupted DOCX file. Unable to read document structure.", status_code=400)

    parts = []

    # 1. Paragraphs
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt:
            parts.append(txt)

    # 2. Tables
    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                # Deduplicate identical adjacent cells (merged cells)
                unique_cells = []
                for c in row_cells:
                    if not unique_cells or c != unique_cells[-1]:
                        unique_cells.append(c)
                parts.append(" | ".join(unique_cells))

    full_text = "\n\n".join(parts)
    return clean_and_normalize_text(full_text)


def extract_from_txt(content: bytes) -> str:
    """Extract clean text from TXT with UTF-8 and Latin-1 fallbacks"""
    try:
        text = content.decode("utf-8")
    except UnicodeDecodeError:
        try:
            text = content.decode("latin-1")
        except Exception as e:
            logger.error(f"Failed to decode TXT: {e}")
            raise ResumeExtractionError("Unable to decode text file. Please ensure it is saved in UTF-8 format.", status_code=400)

    return clean_and_normalize_text(text)


def process_resume_file(filename: str, content: bytes) -> Dict[str, Any]:
    """
    Main processing function for resume upload:
    - Validates size (max 5MB)
    - Detects type via magic bytes
    - Extracts clean text
    - Validates quality (<200 char warning)
    - Saves original file to disk for recruiter access
    """
    # 1. Size Validation
    if len(content) > MAX_FILE_SIZE_BYTES:
        raise ResumeExtractionError(
            f"File is too large ({round(len(content) / (1024 * 1024), 2)} MB). Maximum allowed size is 5 MB.",
            status_code=413
        )

    # 2. File Type Detection
    file_type = detect_file_type(content, filename)

    # 3. Extraction
    if file_type == "pdf":
        extracted_text = extract_from_pdf(content)
    elif file_type == "docx":
        extracted_text = extract_from_docx(content)
    elif file_type == "txt":
        extracted_text = extract_from_txt(content)
    else:
        raise ResumeExtractionError("Unsupported file type.", status_code=400)

    # 4. Save original file to disk
    sanitized_name = re.sub(r'[^a-zA-Z0-9_.-]', '_', filename)
    unique_filename = f"{int(time.time())}_{uuid.uuid4().hex[:6]}_{sanitized_name}"
    saved_path = os.path.join(UPLOAD_DIR, unique_filename)
    try:
        with open(saved_path, "wb") as f:
            f.write(content)
    except Exception as e:
        logger.warning(f"Failed to save original resume to disk: {e}")

    # 5. Quality & Length Checks
    char_count = len(extracted_text)
    word_count = len(extracted_text.split())
    warning = None
    is_valid = True

    if char_count < MIN_VALID_TEXT_LENGTH:
        warning = f"Extracted text is very short ({char_count} characters). A standard resume is usually at least 200 characters. Please review or paste your resume text."
        is_valid = False

    # Check for garbled characters
    non_ascii_ctrl = len(re.findall(r'[^\x20-\x7E\n\t\r]', extracted_text))
    if char_count > 0 and (non_ascii_ctrl / char_count) > 0.08:
        warning = "Extracted text appears to contain unusual or garbled characters. Please review the text box."
        is_valid = False

    return {
        "success": True,
        "filename": filename,
        "file_size": len(content),
        "file_type": file_type,
        "file_path": os.path.relpath(saved_path, os.path.dirname(os.path.dirname(__file__))).replace("\\", "/"),
        "extracted_text": extracted_text,
        "char_count": char_count,
        "word_count": word_count,
        "warning": warning,
        "is_valid": is_valid
    }
